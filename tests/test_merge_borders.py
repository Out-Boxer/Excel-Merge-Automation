import queue
import tempfile
import unittest
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Side

import merge_logic


class BorderCacheTests(unittest.TestCase):
    def setUp(self):
        self.workbook = openpyxl.Workbook()
        self.addCleanup(self.workbook.close)
        self.sheet = self.workbook.active

    def test_repeated_full_border_is_reused_without_changing_style(self):
        side = Side(style="thin", color="FF112233")
        border = Border(left=side, right=side, top=side, bottom=side)
        self.sheet["A1"].border = border
        self.sheet["A2"].border = border
        cache = {}

        first = merge_logic.get_cached_border(self.sheet["A1"], cache)
        second = merge_logic.get_cached_border(self.sheet["A2"], cache)

        self.assertIs(first, second)
        self.assertEqual(first, border)
        self.assertEqual(copy(self.sheet["A1"].border), border)

    def test_incomplete_borders_are_copied_without_caching(self):
        borders = (
            Border(),
            Border(bottom=Side(style="double", color="FFFF0000")),
            Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(color="FFFF0000"),
            ),
        )
        for border in borders:
            with self.subTest(border=border):
                self.sheet["A1"].border = border
                cache = {}

                first = merge_logic.get_cached_border(self.sheet["A1"], cache)
                second = merge_logic.get_cached_border(self.sheet["A1"], cache)

                self.assertEqual(first, border)
                self.assertEqual(second, border)
                self.assertIsNot(first, second)
                self.assertEqual(cache, {})

    def test_merge_preserves_different_borders_from_separate_workbooks(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = []
            expected_borders = []
            border_ids = []
            partial = Border(bottom=Side(style="double", color="FF009900"))

            for index, (style, color) in enumerate(
                (("thin", "FFFF0000"), ("thick", "FF0000FF"))
            ):
                workbook = openpyxl.Workbook()
                try:
                    sheet = workbook.active
                    sheet.title = "Data"
                    side = Side(style=style, color=color)
                    border = Border(left=side, right=side, top=side, bottom=side)
                    for row in (1, 2):
                        sheet.cell(row, 1, f"file{index}-row{row}").border = border
                    sheet["B1"] = "partial"
                    sheet["B1"].border = partial
                    border_ids.append(sheet["A1"]._style.borderId)
                    expected_borders.append(border)
                    path = Path(directory) / f"input{index}.xlsx"
                    workbook.save(path)
                    paths.append(str(path))
                finally:
                    workbook.close()

            # Equal IDs in different files must not share cached border objects.
            self.assertEqual(border_ids[0], border_ids[1])
            output = Path(directory) / "merged.xlsx"
            messages = queue.Queue()
            merge_logic.merge_excel_files(str(output), paths, messages)
            commands = []
            while not messages.empty():
                commands.append(messages.get_nowait()[0])
            self.assertNotIn("show_error", commands)
            self.assertIn("show_info", commands)
            self.assertEqual(commands[-1], "task_done")

            result = openpyxl.load_workbook(output)
            try:
                sheet = result["Data"]
                self.assertEqual(
                    [sheet.cell(row, 1).value for row in range(1, 5)],
                    ["file0-row1", "file0-row2", "file1-row1", "file1-row2"],
                )
                for index, border in enumerate(expected_borders):
                    for row in (index * 2 + 1, index * 2 + 2):
                        self.assertEqual(copy(sheet.cell(row, 1).border), border)
                    self.assertEqual(copy(sheet.cell(index * 2 + 1, 2).border), partial)
            finally:
                result.close()


if __name__ == "__main__":
    unittest.main()
